"""
Export module — writes analysis results to CSV, XLSX, and HTML formats.
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import List, Dict, Any

from .analyzer import AnalysisResult
from .qc import QCResult

logger = logging.getLogger(__name__)


def export_results(
    results: List[AnalysisResult],
    qc_results: List[QCResult],
    parse_errors: List[Dict[str, str]],
    output_dir: Path,
) -> None:
    """
    Export all results to CSV, XLSX, and HTML.

    Args:
        results: List of analysis results.
        qc_results: List of QC results (same order as results).
        parse_errors: List of parse errors encountered.
        output_dir: Directory to write output files.
    """
    output_dir.mkdir(parents=True, exist_ok=True)

    _export_csv(results, qc_results, output_dir / "fasta_analysis.csv")
    _export_invalid_csv(qc_results, results, output_dir / "invalid_sequences.csv")

    try:
        _export_xlsx(results, qc_results, parse_errors, output_dir / "fasta_analysis.xlsx")
    except ImportError:
        logger.warning("openpyxl not available, skipping XLSX export")

    _export_html(results, qc_results, output_dir / "fasta_qc_report.html")

    logger.info("Exported all results to %s", output_dir)


def _export_csv(
    results: List[AnalysisResult],
    qc_results: List[QCResult],
    path: Path,
) -> None:
    """Export primary results to CSV."""
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "Sequence_ID", "Source_File", "Type", "Length",
            "GC_Percent", "AT_Percent", "Ambiguous_Count",
            "Molecular_Weight_Da", "Estimated_pI", "Valid",
            "QC_Status", "QC_Flags",
        ])
        for res, qc in zip(results, qc_results):
            writer.writerow([
                res.record_id,
                res.source_file,
                res.sequence_type,
                res.length,
                res.gc_percent,
                res.at_percent,
                res.ambiguous_count,
                res.molecular_weight,
                res.estimated_pi if res.sequence_type == "PROTEIN" else "",
                res.valid,
                qc.status,
                "; ".join(qc.flags),
            ])
    logger.debug("CSV written: %s", path)


def _export_invalid_csv(
    qc_results: List[QCResult],
    results: List[AnalysisResult],
    path: Path,
) -> None:
    """Export invalid/review sequences to a separate CSV."""
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Sequence_ID", "QC_Status", "Flags", "Error"])
        for qc, res in zip(qc_results, results):
            if qc.status in ("INVALID", "REVIEW"):
                writer.writerow([
                    res.record_id, qc.status, "; ".join(qc.flags), res.error,
                ])


def _export_xlsx(
    results: List[AnalysisResult],
    qc_results: List[QCResult],
    parse_errors: List[Dict[str, str]],
    path: Path,
) -> None:
    """Export results to a multi-sheet XLSX workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    summary_rows = [
        ("Metric", "Value"),
        ("Total sequences", len(results)),
        ("PASS", sum(1 for q in qc_results if q.status == "PASS")),
        ("REVIEW", sum(1 for q in qc_results if q.status == "REVIEW")),
        ("INVALID", sum(1 for q in qc_results if q.status == "INVALID")),
        ("Parse errors", len(parse_errors)),
    ]
    for row in summary_rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = header_font

    # Sheet 2: Sequence Results
    ws2 = wb.create_sheet("Sequence Results")
    headers = ["Sequence_ID", "Source_File", "Type", "Length", "GC%", "AT%",
               "Ambiguous", "MW_Da", "pI", "QC_Status"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.font = header_font
    for res, qc in zip(results, qc_results):
        ws2.append([
            res.record_id, res.source_file, res.sequence_type, res.length,
            res.gc_percent, res.at_percent, res.ambiguous_count,
            res.molecular_weight, res.estimated_pi, qc.status,
        ])

    # Sheet 3: Methods
    ws3 = wb.create_sheet("Methods")
    methods = [
        ("Method", "Description"),
        ("GC%", "(G+C) / (A+T+G+C) × 100 — unambiguous bases only"),
        ("Molecular Weight", "Single-stranded polymer approximation (Biopython-compatible)"),
        ("pI", "Bjellqvist method, simplified bisection solver"),
        ("Reverse Complement", "Standard DNA complement with 5'→3' reversal"),
        ("Source", "BioSyn FASTA Batch Analyzer v1.0.0"),
        ("Website", "https://biosyn-inc.com/"),
    ]
    for row in methods:
        ws3.append(row)

    wb.save(path)
    logger.debug("XLSX written: %s", path)


def _export_html(
    results: List[AnalysisResult],
    qc_results: List[QCResult],
    path: Path,
) -> None:
    """Export a standalone HTML QC report."""
    pass_count = sum(1 for q in qc_results if q.status == "PASS")
    review_count = sum(1 for q in qc_results if q.status == "REVIEW")
    invalid_count = sum(1 for q in qc_results if q.status == "INVALID")

    rows_html = ""
    for res, qc in zip(results, qc_results):
        color = {"PASS": "#28a745", "REVIEW": "#ffc107", "INVALID": "#dc3545"}.get(qc.status, "#6c757d")
        rows_html += f"""
        <tr>
            <td>{res.record_id}</td>
            <td>{res.sequence_type}</td>
            <td>{res.length}</td>
            <td>{res.gc_percent:.1f}%</td>
            <td>{res.ambiguous_count}</td>
            <td style="color:{color};font-weight:bold">{qc.status}</td>
            <td>{'; '.join(qc.flags)}</td>
        </tr>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>BioSyn FASTA Batch QC Report</title>
<style>
body {{ font-family: -apple-system, sans-serif; margin: 2rem; background: #f8f9fa; }}
h1 {{ color: #2c3e50; }}
.stats {{ display: flex; gap: 1rem; margin: 1rem 0; }}
.stat-card {{ padding: 1rem 2rem; border-radius: 8px; color: white; }}
.pass {{ background: #28a745; }}
.review {{ background: #ffc107; color: #333; }}
.invalid {{ background: #dc3545; }}
table {{ border-collapse: collapse; width: 100%; background: white; }}
th, td {{ padding: 0.5rem 1rem; border: 1px solid #dee2e6; text-align: left; }}
th {{ background: #e9ecef; }}
.footer {{ margin-top: 2rem; font-size: 0.85rem; color: #6c757d; }}
</style>
</head>
<body>
<h1>BioSyn FASTA Batch QC Report</h1>
<p>Generated by <a href="https://biosyn-inc.com/">BioSyn Inc.</a> — local-first sequence analysis tools.</p>
<div class="stats">
  <div class="stat-card pass"><strong>PASS</strong><br>{pass_count}</div>
  <div class="stat-card review"><strong>REVIEW</strong><br>{review_count}</div>
  <div class="stat-card invalid"><strong>INVALID</strong><br>{invalid_count}</div>
</div>
<h2>Sequence Results</h2>
<table>
<tr><th>ID</th><th>Type</th><th>Length</th><th>GC%</th><th>Ambiguous</th><th>QC</th><th>Flags</th></tr>
{rows_html}
</table>
<div class="footer">
  <p>Estimated values — molecular weight and pI are approximations. Verify against laboratory standards.</p>
  <p>Learn more at <a href="https://biosyn-inc.com/tools/fasta-sequence-analyzer">BioSyn FASTA Sequence Analyzer</a></p>
</div>
</body>
</html>"""

    path.write_text(html, encoding="utf-8")
    logger.debug("HTML report written: %s", path)
